"""讀取 Revit 明細表匯出檔與文件表格。

Revit「匯出明細表」產生的 .txt 是 Tab 分隔，第一列通常是明細表名稱，
接著才是欄位標題，中間可能夾著群組標題或空白列。所以這裡不假設標題在第一列，
而是往下找「含有編號欄位名稱」的那一列當標題。
"""

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

HEADER_SEARCH_ROWS = 30


class ReadError(Exception):
    pass


@dataclass
class Record:
    key: str
    values: dict
    source: str


@dataclass
class Table:
    path: str
    headers: list
    records: list = field(default_factory=list)


def _decode(raw):
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        # 舊版 Revit／Excel 在繁中 Windows 常存成 Big5
        return raw.decode("cp950")


def _read_text_rows(path):
    text = _decode(Path(path).read_bytes())
    sample = text[:4096]
    delimiter = "\t" if sample.count("\t") >= sample.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _read_excel_rows(path, sheet):
    try:
        import openpyxl
    except ImportError as exc:
        raise ReadError("讀取 Excel 需要 openpyxl：pip install openpyxl") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise ReadError(f"{path} 找不到工作表「{sheet}」，現有：{workbook.sheetnames}")
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.worksheets[0]
        rows = [["" if cell is None else cell for cell in row] for row in worksheet.iter_rows(values_only=True)]
        return rows, worksheet.title
    finally:
        workbook.close()


def _clean_header(value):
    return str(value).strip() if value is not None else ""


def read_table(path, key_column, sheet=None):
    """讀檔並回傳 Table；每筆 Record 記下來源位置（檔名:列號）方便回頭查。"""
    path = str(path)
    suffix = Path(path).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        rows, sheet_name = _read_excel_rows(path, sheet)
        location = f"{Path(path).name}!{sheet_name}"
    elif suffix in (".txt", ".csv", ".tsv"):
        rows = _read_text_rows(path)
        location = Path(path).name
    else:
        raise ReadError(f"不支援的檔案格式：{path}（支援 .txt .csv .tsv .xlsx .xlsm）")

    header_index = None
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        if key_column in (_clean_header(cell) for cell in row):
            header_index = index
            break
    if header_index is None:
        raise ReadError(f"{path} 前 {HEADER_SEARCH_ROWS} 列找不到編號欄位「{key_column}」")

    headers = [_clean_header(cell) for cell in rows[header_index]]
    key_position = headers.index(key_column)
    table = Table(path=path, headers=headers)

    for index, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if key_position >= len(row):
            continue
        key = row[key_position]
        # 空編號的列是群組標題、小計或空白列，跳過
        if key is None or str(key).strip() == "":
            continue
        values = {name: (row[i] if i < len(row) else "") for i, name in enumerate(headers) if name}
        table.records.append(Record(key=str(key).strip(), values=values, source=f"{location}:{index}"))
    return table
